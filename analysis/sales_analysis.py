#!/usr/bin/env python3
"""Анализ продаж МИРСТРОЙМАРКЕТ по двум источникам.

Источники:
  * `ОТЧЕТ ПО ОРГАНИЗАЦИИ мир.xlsx`  — финансовые итоги по кварталам (выручка, НДС,
    себестоимость, затраты, налоги);
  * `analysis/vedomost_ostatki.json` — построчно распознанная «Ведомость по остаткам»
    (PDF, период 01.10.2025–19.10.2026, склад Кореличи, группа «2 краски растворители»).

Результаты: analysis/metrics.json, analysis/finance_quarters.csv,
            analysis/otchet-prodazhi.html (автономный дашборд без внешних зависимостей).

Запуск: python3 analysis/sales_analysis.py
"""
from __future__ import annotations

import csv
import json
import os
from typing import Any

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "ОТЧЕТ ПО ОРГАНИЗАЦИИ мир.xlsx")
STOCK_JSON = os.path.join(ROOT, "analysis", "vedomost_ostatki.json")
OUT_METRICS = os.path.join(ROOT, "analysis", "metrics.json")
OUT_FIN_CSV = os.path.join(ROOT, "analysis", "finance_quarters.csv")
OUT_HTML = os.path.join(ROOT, "analysis", "otchet-prodazhi.html")

MONTHS = ["Октябрь 2025", "Ноябрь 2025", "Декабрь 2025", "Январь 2026", "Февраль 2026",
          "Март 2026", "Апрель 2026", "Май 2026", "Июнь 2026", "Июль 2026",
          "Август 2026", "Сентябрь 2026", "Октябрь 2026"]
QUARTERS = ["1кв", "2кв", "3кв", "4кв", "Итого"]

# строки финансового отчёта (нумерация Excel)
ROW = {
    "vat_payable": 4,        # Налог НДС (к уплате)
    "profit_tax": 5,         # Налог на прибыль
    "revenue_kkm": 6,        # Выручка по отчётам ККМ
    "revenue_inv": 7,        # Реализация (накладные)
    "revenue_total": 8,      # Итого выручка (с НДС)
    "vat_sales": 11,         # Итого НДС с выручки
    "goods_no_vat": 12,      # Товары, реализация без НДС
    "vat_input": 14,         # Входящий НДС
    "vat_input_net": 15,     # Входящий НДС с вычетом
    "cogs": 16,              # Себестоимость реализации
    "opex": 17,              # Себестоимость затрат (сумма статей 18..53)
}
EXPENSE_ROWS = range(18, 54)


def read_finance() -> dict[str, Any]:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Лист1"]
    cols = {q: chr(ord("D") + i) for i, q in enumerate(QUARTERS)}  # D..H

    def val(row: int, q: str) -> float:
        v = ws[f"{cols[q]}{row}"].value
        return float(v) if isinstance(v, (int, float)) else 0.0

    fin: dict[str, Any] = {"quarters": {}, "expenses": []}
    for q in QUARTERS:
        d = {k: val(r, q) for k, r in ROW.items()}
        d["revenue_net"] = d["revenue_total"] - d["vat_sales"]
        d["gross_profit"] = d["revenue_net"] - d["cogs"]
        d["gross_margin"] = (d["gross_profit"] / d["revenue_net"] * 100) if d["revenue_net"] else 0.0
        d["op_profit"] = d["gross_profit"] - d["opex"]
        d["op_margin"] = (d["op_profit"] / d["revenue_net"] * 100) if d["revenue_net"] else 0.0
        d["net_profit"] = d["op_profit"] - d["profit_tax"]
        d["net_margin"] = (d["net_profit"] / d["revenue_net"] * 100) if d["revenue_net"] else 0.0
        fin["quarters"][q] = d

    for r in EXPENSE_ROWS:
        name = ws[f"B{r}"].value
        if not name:
            continue
        total = val(r, "Итого")
        if total:
            fin["expenses"].append({
                "name": str(name).strip(),
                "q": [val(r, q) for q in QUARTERS[:4]],
                "total": total,
            })
    fin["expenses"].sort(key=lambda e: -e["total"])
    return fin


def read_stock() -> list[dict[str, Any]]:
    with open(STOCK_JSON, encoding="utf-8") as fh:
        return json.load(fh)


def stock_metrics(rows: list[dict[str, Any]]) -> dict[str, Any]:
    s = lambda k: sum((r.get(k) or 0) for r in rows)
    monthly = [{"month": m, "units": s(m)} for m in MONTHS]
    deficit = sorted(
        (r for r in rows if r["konechny_ostatok"] < r["min_rezerv"]),
        key=lambda r: r["min_rezerv"] - r["konechny_ostatok"], reverse=True)
    dead = [r for r in rows if not (r["rashod"] or 0)]
    excess = sorted(
        (r for r in rows if r["konechny_ostatok"] >= 2 * r["min_rezerv"]),
        key=lambda r: r["konechny_ostatok"] - r["min_rezerv"], reverse=True)
    top = sorted(rows, key=lambda r: -(r["rashod"] or 0))
    avg_stock = (s("nach_ostatok") + s("konechny_ostatok")) / 2
    full_months = [m for m in MONTHS[6:12]]            # апрель–сентябрь 2026, полные месяцы
    pace = sum((r[m] or 0) for r in rows for m in full_months) / len(full_months)
    return {
        "sku": len(rows),
        "warehouse": sorted({r["sklad"] for r in rows}),
        "group": sorted({r["gruppa"] for r in rows}),
        "opening": s("nach_ostatok"),
        "receipts": s("prihod"),
        "sold_units": s("rashod"),
        "closing": s("konechny_ostatok"),
        "min_reserve": s("min_rezerv"),
        "monthly": monthly,
        "turns_period": round(s("rashod") / avg_stock, 2) if avg_stock else None,
        "coverage_months": round(s("konechny_ostatok") / pace, 2) if pace else None,
        "monthly_pace_units": round(pace, 1),
        "deficit_count": len(deficit),
        "deficit_zero_stock": sum(1 for r in deficit if not r["konechny_ostatok"]),
        "dead_count": len(dead),
        "excess_count": len(excess),
        "top_sold": [{"name": r["nomenklatura"], "units": r["rashod"],
                      "closing": r["konechny_ostatok"], "reserve": r["min_rezerv"]}
                     for r in top[:10] if r["rashod"]],
        "top10_share": round(sum(r["rashod"] or 0 for r in top[:10]) / s("rashod") * 100, 1),
        "deficit_top": [{"name": r["nomenklatura"], "closing": r["konechny_ostatok"],
                         "reserve": r["min_rezerv"], "sold": r["rashod"]} for r in deficit[:12]],
        "dead": [{"name": r["nomenklatura"], "closing": r["konechny_ostatok"]} for r in dead],
        "excess_top": [{"name": r["nomenklatura"], "closing": r["konechny_ostatok"],
                        "reserve": r["min_rezerv"], "sold": r["rashod"]} for r in excess[:10]],
        "months_sum_vs_rashod": {
            "months_sum": sum(m["units"] for m in monthly),
            "rashod": s("rashod"),
            "rows_equal": sum(1 for r in rows
                              if abs(sum((r[m] or 0) for m in MONTHS) - (r["rashod"] or 0)) < 1e-9),
        },
    }


# ---------------------------------------------------------------- HTML-дашборд
CSS = """
:root{--bg:#0f172a;--card:#1e293b;--line:#334155;--txt:#e2e8f0;--mut:#94a3b8;
--acc:#38bdf8;--good:#34d399;--warn:#fbbf24;--bad:#f87171}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--txt);
font:14px/1.5 system-ui,-apple-system,'Segoe UI',Roboto,sans-serif}
.wrap{max-width:1180px;margin:0 auto;padding:28px 20px 60px}
h1{font-size:26px;margin:0 0 4px}h2{font-size:18px;margin:34px 0 12px;
border-left:4px solid var(--acc);padding-left:10px}
.sub{color:var(--mut);margin:0 0 22px}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}
.kpi{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:14px 16px}
.kpi .v{font-size:22px;font-weight:700;margin-top:2px}
.kpi .l{color:var(--mut);font-size:12px;text-transform:uppercase;letter-spacing:.04em}
.kpi .d{color:var(--mut);font-size:12px;margin-top:4px}
.grid2{display:grid;grid-template-columns:repeat(auto-fit,minmax(430px,1fr));gap:16px}
.card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:16px}
.card h3{margin:0 0 10px;font-size:15px}
table{width:100%;border-collapse:collapse;font-size:13px}
th,td{padding:6px 8px;text-align:left;border-bottom:1px solid var(--line);vertical-align:top}
th{color:var(--mut);font-weight:600;font-size:12px;text-transform:uppercase;letter-spacing:.03em}
td.n,th.n{text-align:right;font-variant-numeric:tabular-nums}
tr.tot td{font-weight:700;border-top:2px solid var(--line)}
.note{background:#172554;border:1px solid #1e40af;border-radius:10px;padding:12px 14px;
color:#bfdbfe;font-size:13px;margin:14px 0}
svg text{fill:var(--txt);font:11px system-ui,sans-serif}
svg .mut{fill:var(--mut)}
.legend{display:flex;gap:16px;flex-wrap:wrap;color:var(--mut);font-size:12px;margin-top:8px}
.legend i{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:5px}
.pos{color:var(--good)}.neg{color:var(--bad)}.warnc{color:var(--warn)}
footer{color:var(--mut);font-size:12px;margin-top:36px;border-top:1px solid var(--line);padding-top:14px}
"""


def fmt(v: float, nd: int = 0) -> str:
    s = f"{v:,.{nd}f}".replace(",", " ")
    return s


def bar_chart(series: list[dict], labels: list[str], h: int = 260,
              ytitle: str = "") -> str:
    """Группированная столбчатая диаграмма: series = [{name,color,values}]."""
    w, pad_l, pad_b, pad_t = 560, 62, 34, 18
    n, m = len(labels), len(series)
    gw = (w - pad_l - 10) / n
    bw = min(30.0, (gw - 14) / m)
    vmax = max(max(s["values"]) for s in series) or 1
    scale = (h - pad_b - pad_t) / vmax
    out = [f'<svg viewBox="0 0 {w} {h}" role="img">']
    for i in range(5):
        y = pad_t + (h - pad_b - pad_t) * i / 4
        v = vmax * (1 - i / 4)
        out.append(f'<line x1="{pad_l}" y1="{y:.1f}" x2="{w-6}" y2="{y:.1f}" stroke="#334155" stroke-width="0.6"/>')
        out.append(f'<text x="{pad_l-6}" y="{y+3:.1f}" text-anchor="end" class="mut">{fmt(v)}</text>')
    for gi, lab in enumerate(labels):
        x0 = pad_l + gi * gw + (gw - bw * m) / 2
        for si, s in enumerate(series):
            v = s["values"][gi]
            bh = v * scale
            x = x0 + si * bw
            y = h - pad_b - bh
            out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{bw-2:.1f}" height="{bh:.1f}" '
                       f'fill="{s["color"]}" rx="2"><title>{lab} · {s["name"]}: {fmt(v)}</title></rect>')
        out.append(f'<text x="{pad_l + gi*gw + gw/2:.1f}" y="{h-pad_b+16}" '
                   f'text-anchor="middle" class="mut">{lab}</text>')
    out.append(f'<text x="{pad_l}" y="{12}" class="mut">{ytitle}</text>')
    out.append("</svg>")
    return "".join(out)


def hbar_chart(items: list[tuple[str, float]], color: str, h: int | None = None,
               unit: str = "") -> str:
    w, pad_l = 560, 250
    h = h or (len(items) * 26 + 26)
    vmax = max(v for _, v in items) or 1
    out = [f'<svg viewBox="0 0 {w} {h}" role="img">']
    for i, (name, v) in enumerate(items):
        y = 14 + i * 26
        bw = (w - pad_l - 60) * v / vmax
        label = name if len(name) <= 38 else name[:37] + "…"
        out.append(f'<text x="{pad_l-8}" y="{y+11}" text-anchor="end">{label}</text>')
        out.append(f'<rect x="{pad_l}" y="{y}" width="{max(bw,1):.1f}" height="16" fill="{color}" rx="2">'
                   f'<title>{name}: {fmt(v)}</title></rect>')
        out.append(f'<text x="{pad_l+bw+6:.1f}" y="{y+12}" class="mut">{fmt(v)}{unit}</text>')
    out.append("</svg>")
    return "".join(out)


def build_html(fin: dict, st: dict) -> str:
    q = fin["quarters"]
    qs = ["1кв", "2кв", "3кв"]
    t = q["Итого"]
    kpis = [
        ("Выручка (с НДС)", fmt(t["revenue_total"]), f"ККМ {fmt(t['revenue_kkm'])} + накладные {fmt(t['revenue_inv'])}"),
        ("Выручка без НДС", fmt(t["revenue_net"]), f"НДС с выручки {fmt(t['vat_sales'])}"),
        ("Валовая прибыль", fmt(t["gross_profit"]), f"маржа {t['gross_margin']:.1f}%"),
        ("Операционная прибыль", fmt(t["op_profit"]), f"маржа {t['op_margin']:.1f}%"),
        ("Чистая прибыль", fmt(t["net_profit"]), f"маржа {t['net_margin']:.1f}%, налог {fmt(t['profit_tax'])}"),
        ("НДС к уплате", fmt(t["vat_payable"]), f"входящий НДС {fmt(t['vat_input_net'])}"),
        ("Продано, шт", fmt(st["sold_units"]), f"склад Кореличи, {st['sku']} SKU"),
        ("Оборачиваемость запаса", f"{st['turns_period']}×", f"покрытие остатком {st['coverage_months']} мес"),
        ("Дефицит SKU", str(st["deficit_count"]), f"из них с нулевым остатком {st['deficit_zero_stock']}"),
    ]
    html = [f"<!doctype html><html lang='ru'><head><meta charset='utf-8'>"
            f"<meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<title>Анализ продаж — МИРСТРОЙМАРКЕТ</title><style>{CSS}</style></head><body><div class='wrap'>"]
    html.append("<h1>Анализ продаж — МИРСТРОЙМАРКЕТ</h1>")
    html.append("<p class='sub'>Финансовый отчёт: 1–3 кварталы 2026 (9 месяцев). "
                "Товарный отчёт: ведомость по остаткам, 01.10.2025–19.10.2026, склад Кореличи, "
                "группа «2 краски растворители». Все суммы — в рублях отчёта.</p>")
    html.append("<div class='kpis'>")
    for l, v, d in kpis:
        html.append(f"<div class='kpi'><div class='l'>{l}</div><div class='v'>{v}</div><div class='d'>{d}</div></div>")
    html.append("</div>")

    html.append("<h2>1. Финансовый результат по кварталам</h2><div class='grid2'><div class='card'>")
    html.append("<h3>Выручка, себестоимость и прибыль</h3>")
    html.append(bar_chart([
        {"name": "Выручка без НДС", "color": "#38bdf8", "values": [q[k]["revenue_net"] for k in qs]},
        {"name": "Себестоимость товаров", "color": "#64748b", "values": [q[k]["cogs"] for k in qs]},
        {"name": "Валовая прибыль", "color": "#34d399", "values": [q[k]["gross_profit"] for k in qs]},
        {"name": "Затраты (opex)", "color": "#fbbf24", "values": [q[k]["opex"] for k in qs]},
        {"name": "Чистая прибыль", "color": "#a78bfa", "values": [q[k]["net_profit"] for k in qs]},
    ], qs, ytitle="руб."))
    html.append("<div class='legend'>" + "".join(
        f"<span><i style='background:{s['color']}'></i>{s['name']}</span>" for s in [
            {"name": "Выручка без НДС", "color": "#38bdf8"}, {"name": "Себестоимость товаров", "color": "#64748b"},
            {"name": "Валовая прибыль", "color": "#34d399"}, {"name": "Затраты (opex)", "color": "#fbbf24"},
            {"name": "Чистая прибыль", "color": "#a78bfa"}]) + "</div></div>")
    html.append("<div class='card'><h3>Квартальная таблица</h3><table><tr><th>Показатель</th>"
                + "".join(f"<th class='n'>{k}</th>" for k in QUARTERS) + "</tr>")
    lines = [
        ("Выручка ККМ (розница)", "revenue_kkm"), ("Выручка накладные (опт)", "revenue_inv"),
        ("Итого выручка с НДС", "revenue_total"), ("НДС с выручки", "vat_sales"),
        ("Выручка без НДС", "revenue_net"), ("Себестоимость реализации", "cogs"),
        ("Валовая прибыль", "gross_profit"), ("Затраты (opex)", "opex"),
        ("Операционная прибыль", "op_profit"), ("Налог на прибыль 20%", "profit_tax"),
        ("Чистая прибыль", "net_profit"), ("НДС к уплате", "vat_payable"),
    ]
    for name, key in lines:
        html.append(f"<tr><td>{name}</td>" + "".join(
            f"<td class='n'>{fmt(q[k][key])}</td>" for k in QUARTERS) + "</tr>")
    html.append("<tr class='tot'><td>Маржа валовая, %</td>" + "".join(
        f"<td class='n'>{q[k]['gross_margin']:.1f}</td>" for k in QUARTERS) + "</tr>")
    html.append("<tr class='tot'><td>Маржа чистая, %</td>" + "".join(
        f"<td class='n'>{q[k]['net_margin']:.1f}</td>" for k in QUARTERS) + "</tr>")
    html.append("</table></div></div>")

    html.append("<h2>2. Продажи в натуральном выражении (шт)</h2>")
    html.append("<div class='note'>Помесячные колонки ведомости используются как индикатор динамики: "
                f"их сумма ({fmt(st['months_sum_vs_rashod']['months_sum'])} шт) превышает колонку «Расход» "
                f"({fmt(st['months_sum_vs_rashod']['rashod'])} шт) по 74 из 111 позиций, поэтому итоговые "
                "продажи за период считаются по колонке «Расход», проверенной балансом "
                "«начальный остаток + приход − расход = конечный остаток» (сходится во всех 111 строках).</div>")
    html.append("<div class='grid2'><div class='card'><h3>Динамика по месяцам (помесячные колонки ведомости)</h3>")
    html.append(bar_chart([{"name": "шт", "color": "#38bdf8",
                            "values": [m["units"] for m in st["monthly"]]}],
                          ["10.25", "11.25", "12.25", "01.26", "02.26", "03.26", "04.26",
                           "05.26", "06.26", "07.26", "08.26", "09.26", "10.26*"],
                          ytitle="шт; *октябрь 2026 — неполный месяц (по 19.10)"))
    html.append("</div><div class='card'><h3>Топ-10 позиций по продажам за период</h3>")
    html.append(hbar_chart([(i["name"], i["units"]) for i in st["top_sold"]], "#34d399", unit=" шт"))
    html.append(f"<div class='d' style='color:var(--mut);font-size:12px;margin-top:6px'>"
                f"Топ-10 позиций дают {st['top10_share']}% всех продаж в штуках.</div></div></div>")

    html.append("<h2>3. Структура затрат за 9 месяцев</h2><div class='card'>")
    html.append(hbar_chart([(e["name"], e["total"]) for e in fin["expenses"][:10]], "#fbbf24"))
    tot_opex = t["opex"]
    html.append(f"<div style='color:var(--mut);font-size:12px;margin-top:8px'>Всего затрат "
                f"{fmt(tot_opex)} руб. = {tot_opex/t['revenue_net']*100:.1f}% выручки без НДС. "
                f"ФОТ с отчислениями (фонд з/пл + ФСЗН) — "
                f"{(fin['expenses'][0]['total']+fin['expenses'][1]['total'])/tot_opex*100:.0f}% затрат.</div></div>")

    html.append("<h2>4. Состояние склада на 19.10.2026</h2><div class='grid2'>")
    html.append(f"<div class='card'><h3>Дефицит: остаток ниже минимального резерва "
                f"({st['deficit_count']} SKU)</h3><table><tr><th>Позиция</th><th class='n'>Остаток</th>"
                "<th class='n'>Резерв</th><th class='n'>Продано</th></tr>")
    for d in st["deficit_top"]:
        cls = "neg" if not d["closing"] else "warnc"
        html.append(f"<tr><td>{d['name']}</td><td class='n {cls}'>{fmt(d['closing'])}</td>"
                    f"<td class='n'>{fmt(d['reserve'])}</td><td class='n'>{fmt(d['sold'])}</td></tr>")
    html.append("</table></div>")
    html.append(f"<div class='card'><h3>Излишки и неликвиды</h3><table><tr><th>Позиция</th>"
                "<th class='n'>Остаток</th><th class='n'>Резерв</th><th class='n'>Продано</th></tr>")
    for d in st["excess_top"]:
        html.append(f"<tr><td>{d['name']}</td><td class='n'>{fmt(d['closing'])}</td>"
                    f"<td class='n'>{fmt(d['reserve'])}</td><td class='n'>{fmt(d['sold'])}</td></tr>")
    html.append("<tr class='tot'><td colspan='4'>Неликвиды (продаж не было): "
                + "; ".join(f"{d['name']} — {fmt(d['closing'])} шт" for d in st["dead"])
                + "</td></tr></table></div></div>")

    html.append("<footer>Источники: «ОТЧЕТ ПО ОРГАНИЗАЦИИ мир.xlsx» (квартальные итоги с формулами) и "
                "«Ведомость мав.pdf» (ведомость по остаткам, 111 строк, распознано по координатам). "
                "Отчёт управленческий и не заменяет бухгалтерскую отчётность. "
                "Сформировано: analysis/sales_analysis.py.</footer>")
    html.append("</div></body></html>")
    return "".join(html)


def main() -> None:
    fin = read_finance()
    rows = read_stock()
    st = stock_metrics(rows)

    with open(OUT_FIN_CSV, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=";")
        keys = ["revenue_kkm", "revenue_inv", "revenue_total", "vat_sales", "revenue_net",
                "cogs", "gross_profit", "gross_margin", "opex", "op_profit", "op_margin",
                "profit_tax", "net_profit", "net_margin", "vat_payable", "vat_input_net"]
        w.writerow(["показатель"] + QUARTERS)
        for k in keys:
            w.writerow([k] + [round(fin["quarters"][q][k], 2) for q in QUARTERS])
        w.writerow([])
        w.writerow(["статья затрат"] + QUARTERS[:4] + ["итого"])
        for e in fin["expenses"]:
            w.writerow([e["name"]] + [round(x, 2) for x in e["q"]] + [round(e["total"], 2)])

    metrics = {"finance": fin, "stock": st}
    with open(OUT_METRICS, "w", encoding="utf-8") as fh:
        json.dump(metrics, fh, ensure_ascii=False, indent=1)

    with open(OUT_HTML, "w", encoding="utf-8") as fh:
        fh.write(build_html(fin, st))

    t = fin["quarters"]["Итого"]
    print("финансы (итого): выручка с НДС", round(t["revenue_total"], 2),
          "| без НДС", round(t["revenue_net"], 2),
          "| валовая", round(t["gross_profit"], 2), f"({t['gross_margin']:.1f}%)",
          "| операционная", round(t["op_profit"], 2),
          "| чистая", round(t["net_profit"], 2), f"({t['net_margin']:.1f}%)")
    print("склад: SKU", st["sku"], "| продано шт", st["sold_units"],
          "| оборотов", st["turns_period"], "| покрытие", st["coverage_months"], "мес",
          "| дефицит", st["deficit_count"], "| неликвид", st["dead_count"])
    print("записано:", OUT_FIN_CSV, OUT_METRICS, OUT_HTML, sep="\n  ")


if __name__ == "__main__":
    main()
